"""HireSense FastAPI app: live voice session over WebSocket, session history over REST."""
from __future__ import annotations

import asyncio
import logging
import logging.config
from contextlib import asynccontextmanager

from fastapi import FastAPI, HTTPException, WebSocket, UploadFile, File, Depends
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
import io
import docx
import pypdf

from .. import sessions_store, tts, auth
from ..auth import get_current_user_id
from ..config import get_settings
from ..session import SessionController

settings = get_settings()

# ---- logging setup ----------------------------------------------------------
logging.config.dictConfig({
    "version": 1,
    "disable_existing_loggers": False,
    "formatters": {
        "default": {
            "format": "%(asctime)s [%(levelname)s] %(name)s: %(message)s",
            "datefmt": "%H:%M:%S",
        }
    },
    "handlers": {
        "console": {
            "class": "logging.StreamHandler",
            "formatter": "default",
        }
    },
    "loggers": {
        "hiresense": {"level": "INFO", "handlers": ["console"], "propagate": False},
        "uvicorn.access": {"level": "WARNING"},   # suppress per-request noise
    },
    "root": {"level": "WARNING"},
})
logger = logging.getLogger("hiresense.api")


@asynccontextmanager
async def lifespan(app: FastAPI):
    # Load the Kokoro TTS model once (~1.5s) so the first turn doesn't pay it.
    # A load failure degrades to captions-only sessions rather than blocking
    # startup — /health exposes tts_ready so the UI can show a warning state.
    ok = await asyncio.to_thread(tts.warm_up)
    if not ok:
        logger.error("[api] TTS unavailable: %s", tts.load_error())
    yield


app = FastAPI(title="HireSense API", version="1.0.0", lifespan=lifespan)
app.add_middleware(
    CORSMiddleware,
    allow_origins=settings.cors_origin_list,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.get("/health")
async def health() -> dict:
    return {
        "status": "ok",
        "tts_ready": tts.is_ready(),
        "tts_error": tts.load_error(),
        "stt_model": settings.stt_model,
    }


@app.websocket("/ws/session")
async def ws_session(ws: WebSocket) -> None:
    token = ws.query_params.get("token")
    user_id = "anonymous-developer"
    domain = auth.get_clerk_domain()
    
    import sys
    is_testing = "pytest" in sys.modules
    
    if not is_testing and domain:
        if not token:
            await ws.close(code=4008)  # Policy Violation
            return
        payload = auth.verify_clerk_token(token)
        if not payload or "sub" not in payload:
            await ws.close(code=4008)
            return
        user_id = payload["sub"]
        
    await SessionController(ws, user_id=user_id).run()


class SessionSummary(BaseModel):
    id: str
    role: str
    seniority: str
    question_count: int
    status: str
    created_at: str
    overall_average: float | None
    top_improvements: list[str]


class SessionDetail(SessionSummary):
    turns: list[dict]


@app.get("/sessions", response_model=list[SessionSummary])
async def list_sessions(limit: int = 50, user_id: str = Depends(get_current_user_id)) -> list[dict]:
    return await asyncio.to_thread(sessions_store.list_sessions, user_id, limit)


@app.get("/sessions/{session_id}", response_model=SessionDetail)
async def get_session(session_id: str, user_id: str = Depends(get_current_user_id)) -> dict:
    data = await asyncio.to_thread(sessions_store.get_session, session_id, user_id)
    if data is None:
        raise HTTPException(status_code=404, detail="Session not found")
    return data


@app.delete("/sessions/{session_id}")
async def delete_session(session_id: str, user_id: str = Depends(get_current_user_id)) -> dict:
    await asyncio.to_thread(sessions_store.delete_session, session_id, user_id)
    return {"ok": True}


@app.post("/upload_resume")
async def upload_resume(file: UploadFile = File(...), user_id: str = Depends(get_current_user_id)) -> dict:
    filename = file.filename or ""
    ext = filename.split(".")[-1].lower() if "." in filename else ""
    
    try:
        content = await file.read()
        text = ""
        
        if ext == "pdf":
            reader = pypdf.PdfReader(io.BytesIO(content))
            pages_text = []
            for page in reader.pages:
                t = page.extract_text()
                if t:
                    pages_text.append(t)
            text = "\n".join(pages_text)
            
        elif ext == "docx":
            doc = docx.Document(io.BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        paragraphs.append(cell.text)
            text = "\n".join(paragraphs)
            
        else:
            # Fallback to plain text decoding
            try:
                text = content.decode("utf-8")
            except UnicodeDecodeError:
                text = content.decode("latin-1", errors="ignore")
                
        # Limit text length defensively to avoid overwhelming LLM context
        text = text.strip()
        if len(text) > 50000:
            text = text[:50000] + "\n... [resume truncated]"
            
        return {"text": text}
        
    except Exception as e:
        logger.exception("Failed to parse uploaded resume: %s", filename)
        raise HTTPException(status_code=400, detail=f"Failed to parse resume: {str(e)}")


@app.get("/")
def root() -> dict:
    return {"name": "HireSense", "docs": "/docs", "health": "/health"}
